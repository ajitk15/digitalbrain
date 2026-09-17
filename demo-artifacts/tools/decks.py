"""PowerPoint output.

Two shapes. `write_deck` turns a document into slides, which suits an overview
whose level-two headings are already the agenda. `build_slides` takes slide
content written as slides in the first place, for a deck that was never a
document - the release readiness review, for instance, which exists as a deck
and only as a deck.

Both put the speaker's material in the notes rather than on the slide. A slide
carrying a paragraph is a document with worse typography.
"""

from pathlib import Path

from mdparse import Block, front_matter, plain, title_of

NAVY = (0x1F, 0x4E, 0x79)
SLATE = (0x5B, 0x64, 0x70)
AMBER = (0xC9, 0xA2, 0x27)

#: A slide holds this many bullets before the rest moves to the notes.
MAX_BULLETS = 6


def _rgb(triple):
    from pptx.dml.color import RGBColor

    return RGBColor(*triple)


def _new_presentation():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    return presentation


def _title_slide(presentation, title: str, subtitle: str, footer: str) -> None:
    from pptx.util import Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = _rgb(NAVY)

    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.9), Inches(3.9), Inches(11.5), Inches(1.0))
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = subtitle
        run.font.size = Pt(18)
        run.font.color.rgb = _rgb(SLATE)

    box = slide.shapes.add_textbox(Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.6))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = footer
    run.font.size = Pt(11)
    run.font.color.rgb = _rgb(AMBER)


def _content_slide(presentation, heading: str, bullets: list[str], notes: str = "") -> None:
    from pptx.util import Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.8), Inches(1.0))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = heading
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = _rgb(NAVY)

    if bullets:
        box = slide.shapes.add_textbox(Inches(0.95), Inches(1.7), Inches(11.4), Inches(5.0))
        frame = box.text_frame
        frame.word_wrap = True
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = "•  " + bullet
            run.font.size = Pt(16)
            paragraph.space_after = Pt(10)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _table_slide(presentation, heading: str, rows: list[list[str]], notes: str = "") -> None:
    from pptx.util import Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.75), Inches(0.45), Inches(11.8), Inches(0.9))
    frame = box.text_frame
    run = frame.paragraphs[0].add_run()
    run.text = heading
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = _rgb(NAVY)

    shown = rows[:11]
    width = max(len(row) for row in shown)
    shape = slide.shapes.add_table(
        len(shown), width, Inches(0.75), Inches(1.5), Inches(11.8), Inches(0.4 * len(shown))
    )
    table = shape.table
    for r, row in enumerate(shown):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = plain(row[c]) if c < len(row) else ""
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.bold = r == 0
    if notes or len(rows) > 11:
        extra = f"Full table has {len(rows) - 1} rows; {len(shown) - 1} shown."
        slide.notes_slide.notes_text_frame.text = (notes + "\n" + extra).strip()


def write_deck(blocks: list[Block], target: Path, fallback: str) -> None:
    """Turn a document into a deck, one slide per level-two heading."""
    presentation = _new_presentation()
    meta = front_matter(blocks)
    title = title_of(blocks, fallback)
    _title_slide(
        presentation,
        title,
        meta.get("Owner", ""),
        "Demonstration artifact — synthetic content only, not a clinical record.",
    )

    heading = ""
    bullets: list[str] = []
    notes: list[str] = []
    pending_table: list[list[str]] | None = None
    started = False

    def flush() -> None:
        nonlocal heading, bullets, notes, pending_table
        if not heading:
            return
        if pending_table and len(pending_table) > 1:
            _table_slide(presentation, heading, pending_table, "\n\n".join(notes))
        else:
            _content_slide(
                presentation, heading, bullets[:MAX_BULLETS], "\n\n".join(notes + bullets[MAX_BULLETS:])
            )
        heading, bullets, notes, pending_table = "", [], [], None

    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                continue
            if block.level == 2:
                flush()
                heading = plain(block.text)
                started = True
            elif started:
                bullets.append(plain(block.text))
            continue
        if not started:
            continue
        if block.kind == "para":
            text = plain(block.text)
            (bullets if len(text) < 130 else notes).append(text)
        elif block.kind == "list":
            bullets.extend(plain(item) for item in block.items)
        elif block.kind == "table" and any(block.rows[0]):
            if pending_table is None:
                pending_table = block.rows
            else:
                notes.append(" / ".join(plain(cell) for cell in block.rows[0]))
        elif block.kind == "quote":
            notes.append(plain(block.text))
    flush()
    presentation.save(target)


def build_slides(target: Path, title: str, subtitle: str, slides: list[dict]) -> None:
    """Build a deck from slide definitions rather than from a document.

    Each entry is `{"heading": str, "bullets": [str], "rows": [[str]], "notes": str}`
    with `bullets` and `rows` both optional.
    """
    presentation = _new_presentation()
    _title_slide(
        presentation,
        title,
        subtitle,
        "Demonstration artifact — synthetic content only, not a clinical record.",
    )
    for slide in slides:
        if slide.get("rows"):
            _table_slide(presentation, slide["heading"], slide["rows"], slide.get("notes", ""))
        else:
            _content_slide(
                presentation,
                slide["heading"],
                slide.get("bullets", []),
                slide.get("notes", ""),
            )
    presentation.save(target)
